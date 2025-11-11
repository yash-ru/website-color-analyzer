#!/usr/bin/env python3
"""
app.py - Streamlit Web Interface for Domain Screenshot & Color Analysis (Unified Mode)
Usage: streamlit run app.py
"""

import streamlit as st
import os
import asyncio
import pandas as pd
import time
from io import BytesIO
from pathlib import Path
from PIL import Image

# Import custom modules
from screenshotter import DomainScreenshotter
from color_analyzer2 import ColorAnalyzer
from dominant_colors2 import is_neutral

start_time = time.time()

# ---------------- UI SETUP ---------------- #
st.set_page_config(page_title="Website Color Analyzer", layout="centered")
st.title("Website Color Analyzer")

st.markdown("""
    <style>
    .success-box {
        padding: 1rem;
        background-color: #d4edda;
        border-left: 5px solid #28a745;
        border-radius: 5px;
        margin: 1rem 0;
    }
    .error-box {
        padding: 1rem;
        background-color: #f8d7da;
        border-left: 5px solid #dc3545;
        border-radius: 5px;
        margin: 1rem 0;
    }
    </style>
""", unsafe_allow_html=True)

# ---------------- CORE FUNCTIONS ---------------- #
async def capture_screenshot_async(domain, output_folder):
    """Async wrapper for screenshot capture."""
    screenshotter = DomainScreenshotter(output_folder=output_folder, concurrency=5)
    results = await screenshotter.run([domain])
    return results[0] if results else None

def hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))

def text_contrast_color(hex_color):
    """Return black or white text based on hex color luminance."""
    try:
        r, g, b = hex_to_rgb(hex_color)
        luminance = (0.299*r + 0.587*g + 0.114*b) / 255
        return "black" if luminance > 0.6 else "white"
    except:
        return "black"

def get_dominant_colors_from_result(analysis_result, top_n=2):
    """
    Extract top dominant colors and detect theme.
    Returns: dominant_colors list, theme
    """
    colors = [
        {"hex": c["hex"], "rgb": tuple(c["rgb"]) if "rgb" in c else hex_to_rgb(c["hex"]), "pct": c["percentage"]}
        for c in analysis_result["colors"]
    ]

    if not colors:
        return [""]*top_n, "Unknown"

    # Compute weighted average brightness
    avg_brightness = sum(c['pct'] * ((0.299*c['rgb'][0]+0.587*c['rgb'][1]+0.114*c['rgb'][2])/255) for c in colors) / max(sum(c['pct'] for c in colors), 0.0001)
    theme = "Dark" if avg_brightness < 0.5 else "Light"

    # Filter out neutral colors
    non_neutral = [c for c in colors if not is_neutral(c['rgb'])]

    # Choose dominant colors based on theme
    dominant_colors = []
    if theme == "Dark":
        # Bright accents in dark theme
        bright_colors = [c for c in non_neutral if ((0.299*c['rgb'][0]+0.587*c['rgb'][1]+0.114*c['rgb'][2])/255) > 0.15]
        if not bright_colors:
            bright_colors = non_neutral or colors
        bright_colors.sort(key=lambda x: x['pct']*((0.299*x['rgb'][0]+0.587*x['rgb'][1]+0.114*x['rgb'][2])/255), reverse=True)
        dominant_colors = bright_colors[:top_n]
    else:
        # Just pick top percentages in light theme
        non_neutral.sort(key=lambda x: -x['pct'])
        dominant_colors = non_neutral[:top_n]

    # Ensure top_n elements
    while len(dominant_colors) < top_n:
        dominant_colors.append({"hex": ""})

    return [c['hex'] for c in dominant_colors], theme


def analyze_domain(domain, output_folder, analysis_folder, num_colors, visualize):
    """Capture screenshot, analyze colors, and add dominant colors."""
    try:
        # Reuse the bulk logic for consistency
        results = asyncio.run(run_bulk_analysis_async([domain], output_folder, analysis_folder, num_colors, visualize))
        result = results[0] if results else {"domain": domain, "error": "Processing failed"}
        if not result or not result.get("success"):
            return {"domain": domain, "error": result.get("error", "Screenshot failed")}

        analyzer = ColorAnalyzer(output_folder=analysis_folder)
        analysis_result = analyzer.analyze_single(result["output_path"], num_colors=num_colors, visualize=visualize)

        if not analysis_result:
            return {"domain": domain, "error": "Color analysis failed"}

        # Sort top colors by percentage
        colors_sorted = sorted(analysis_result["colors"], key=lambda x: x["percentage"], reverse=True)

        flattened = {"domain": analysis_result["domain"]}
        for i, color in enumerate(colors_sorted, start=1):
            flattened[f"top{i}_hex"] = color["hex"]
            flattened[f"top{i}_pct"] = color["percentage"]

        # Extract dominant colors
        dominant_colors, theme = get_dominant_colors_from_result(analysis_result, top_n=2)
        flattened["dominant_color_1"] = dominant_colors[0] if len(dominant_colors) > 0 else ""
        flattened["dominant_color_2"] = dominant_colors[1] if len(dominant_colors) > 1 else ""
        flattened["theme"] = theme

        return flattened

    except Exception as e:
        return {"domain": domain, "error": str(e)}

async def run_bulk_analysis_async(domains, output_folder, analysis_folder, num_colors, visualize=False):
    """Parallel screenshot capture + parallel color analysis"""
    
    # STEP 1: Capture ALL screenshots in parallel
    st.info(f"📸 Capturing screenshots for {len(domains)} domains...")
    screenshotter = DomainScreenshotter(output_folder=output_folder, concurrency=5)
    screenshot_results = await screenshotter.run(domains)
    
    # STEP 2: Filter successful screenshots
    successful = [r for r in screenshot_results if r.get("success")]
    failed = [r for r in screenshot_results if not r.get("success")]
    
    if failed:
        st.warning(f"⚠️ {len(failed)} screenshots failed: {[r['domain'] for r in failed]}")
    
    if not successful:
        return []
    
    # STEP 3: Analyze ALL colors in parallel using multiprocessing
    st.info(f"🎨 Analyzing colors for {len(successful)} screenshots...")
    analyzer = ColorAnalyzer(screenshots_folder=output_folder, output_folder=analysis_folder)
    
    # Use the existing multiprocessing capability
    analysis_results = analyzer.analyze_all(num_colors=num_colors, visualize=visualize, workers=6)
    
    # STEP 4: Merge results and add dominant colors
    final_results = []
    for analysis in analysis_results:
        # Find corresponding screenshot result
        screenshot_result = next((s for s in successful if s['domain'].replace('https://', '').replace('http://', '').rstrip('/').replace('/', '_').replace(':', '_').replace('www.', '') == analysis['domain']), None)
        
        # Flatten colors
        flattened = {"domain": analysis["domain"]}
        colors_sorted = sorted(analysis["colors"], key=lambda x: x["percentage"], reverse=True)
        
        for i, color in enumerate(colors_sorted, start=1):
            flattened[f"top{i}_hex"] = color["hex"]
            flattened[f"top{i}_pct"] = color["percentage"]
        
        # Add dominant colors
        dominant_colors, theme = get_dominant_colors_from_result(analysis, top_n=2)
        flattened["dominant_color_1"] = dominant_colors[0] if len(dominant_colors) > 0 else ""
        flattened["dominant_color_2"] = dominant_colors[1] if len(dominant_colors) > 1 else ""
        flattened["theme"] = theme
        
        final_results.append(flattened)
    
    return final_results

def style_dataframe(df):
    """Apply styling to hex and percentage columns."""
    hex_cols = [col for col in df.columns if col.startswith("top") and "_hex" in col] + \
               [col for col in df.columns if col.startswith("dominant_color")]

    pct_cols = [col for col in df.columns if col.startswith("top") and "_pct" in col]

    df_styled = df.style.applymap(
        lambda val: f"background-color: {val}; color: {text_contrast_color(val)}",
        subset=hex_cols
    )

    df_styled = df_styled.bar(subset=pct_cols, color="#4caf50", vmin=0, vmax=1)
    return df_styled

def download_excel(df, filename):
    """Download DataFrame as Excel preserving hex background colors."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font
    except ImportError:
        st.error("Please install openpyxl: pip install openpyxl")
        return None

    wb = Workbook()
    ws = wb.active
    ws.append(df.columns.tolist())

    for _, row in df.iterrows():
        ws.append([row[c] for c in df.columns])

    for i, col in enumerate(df.columns, start=1):
        if "_hex" in col or "dominant_color" in col:
            for j, val in enumerate(df[col], start=2):
                try:
                    fill = PatternFill(start_color=val.lstrip("#"), end_color=val.lstrip("#"), fill_type="solid")
                    text_color = text_contrast_color(val)
                    font = Font(color="000000" if text_color=="black" else "FFFFFF")
                    ws.cell(row=j, column=i).fill = fill
                    ws.cell(row=j, column=i).font = font
                except:
                    continue
    with BytesIO() as output:
        wb.save(output)
        return output.getvalue()

# ---------------- MAIN APP ---------------- #
def main():
    st.write("Enter one or more domains (each on a new line)")
    domains_input = st.text_area("Domains", placeholder="example.com\nopenai.com\ngoogle.com", height=150)

    #num_colors = st.slider("Number of Colors", 3, 10, 5)
    #visualize = st.toggle("Generate Color Palette Images", value=True)
    
    num_colors = 5
    visualize = False

    output_folder = "screenshots"
    analysis_folder = "color_analysis"

    if st.button("Analyze", use_container_width=True) and domains_input.strip():
        domains = [d.strip() for d in domains_input.splitlines() if d.strip()]

        # Single domain
        if len(domains) == 1:
            domain = domains[0]
            with st.spinner(f"Processing {domain}..."):
                results = asyncio.run(run_bulk_analysis_async([domain], output_folder, analysis_folder, num_colors, visualize))
                result = results[0] if results else {"domain": domain, "error": "Processing failed"}

            if "error" in result:
                st.error(f"❌ {domain}: {result['error']}")
            else:
                st.success(f"✅ Analysis completed for {domain}")

                screenshot_path = os.path.join(output_folder, f"{result['domain']}.png")
                palette_path = os.path.join(analysis_folder, f"{result['domain']}_palette.png")

                if os.path.exists(screenshot_path):
                    st.image(screenshot_path, caption="Screenshot", use_container_width=True)
                if visualize and os.path.exists(palette_path):
                    st.image(palette_path, caption="Color Palette", use_container_width=True)

                st.subheader("🎨 Color Summary (Top Colors + Dominant)")
                df = pd.DataFrame([result])
                st.dataframe(style_dataframe(df), use_container_width=True)

                excel_bytes = download_excel(df, f"{domain}_colors.xlsx")
                if excel_bytes:
                    st.download_button("📥 Download Excel", excel_bytes, f"{domain}_colors.xlsx",
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        # Bulk domains
        else:
            results = asyncio.run(run_bulk_analysis_async(domains, output_folder, analysis_folder, num_colors, visualize))
            
            valid_results = [r for r in results if "error" not in r]

            valid_results = [r for r in results if "error" not in r]
            error_results = [r for r in results if "error" in r]

            if valid_results:
                st.markdown("---")
                st.subheader("Bulk Analysis Results")

                df_out = pd.DataFrame(valid_results)
                st.dataframe(style_dataframe(df_out), use_container_width=True)

                excel_bytes = download_excel(df_out, "bulk_color_results.xlsx")
                if excel_bytes:
                    st.download_button("📥 Download All Results", excel_bytes,
                                       "bulk_color_results.xlsx",
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

            if error_results:
                st.warning(f"⚠️ {len(error_results)} domains failed.")
                for e in error_results:
                    st.write(f"- {e['domain']}: {e['error']}")


if __name__ == "__main__":
    main()

total_elapsed = time.time() - start_time
st.success(f"✅ Total processing time: {total_elapsed:.2f} seconds")
